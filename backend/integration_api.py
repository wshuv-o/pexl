"""Authenticated FinAi workbook integration; also mountable in the OCR app."""
from __future__ import annotations

import base64
from copy import copy
from datetime import datetime, timezone
from functools import lru_cache
from io import BytesIO
import json
import logging
import os
from pathlib import Path
import secrets
import shutil
import sqlite3
import subprocess
from typing import Literal
from uuid import uuid4
from zipfile import ZipFile

from fastapi import APIRouter, Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.responses import Response
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from pydantic import BaseModel, ConfigDict, Field, model_validator

MAX_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_CELLS = 250_000
MAX_EXPORT_FILES = 1000
MAX_EXPORT_ROWS = 100_000
EXPORT_TIMEOUT_SECONDS = 120
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
security = HTTPBearer(auto_error=False)
log = logging.getLogger(__name__)


def owner(credentials: HTTPAuthorizationCredentials | None = Depends(security)) -> str:
    """Map separately issued secrets to client namespaces; fail closed."""
    try:
        clients = json.loads(os.environ.get('PEXL_API_KEYS', '{}'))
        if not isinstance(clients, dict) or any(
            not isinstance(k, str) or not isinstance(v, str) or len(v) < 32
            for k, v in clients.items()
        ):
            raise ValueError('Invalid API key configuration')
        if len(set(clients.values())) != len(clients):
            raise ValueError('Client keys must be distinct')
    except (ValueError, TypeError):
        raise HTTPException(503, 'Integration authentication is not configured')
    if not clients:
        raise HTTPException(503, 'Integration authentication is not configured')
    token = credentials.credentials if credentials else ''
    for client, key in clients.items():
        if secrets.compare_digest(token.encode(), key.encode()):
            return client
    raise HTTPException(401, 'Invalid bearer token', headers={'WWW-Authenticate': 'Bearer'})


def database():
    root = Path(os.environ.get('PEXL_INTEGRATION_DIR', Path(__file__).parent / 'integration-data'))
    root.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(root / 'integration.sqlite3', timeout=30)
    db.row_factory = sqlite3.Row
    db.execute('''CREATE TABLE IF NOT EXISTS artifacts (
        id TEXT PRIMARY KEY, owner TEXT NOT NULL, kind TEXT NOT NULL,
        metadata TEXT NOT NULL, content BLOB NOT NULL)''')
    return db


def save(client: str, kind: str, metadata: dict, content: bytes = b'') -> dict:
    result = {**metadata, 'id': uuid4().hex, 'kind': kind,
              'created_at': datetime.now(timezone.utc).isoformat()}
    db = database()
    try:
        with db:
            db.execute('INSERT INTO artifacts VALUES (?, ?, ?, ?, ?)',
                       (result['id'], client, kind, json.dumps(result), content))
    finally:
        db.close()
    return result


def fetch(client: str, artifact_id: str, kind: str):
    db = database()
    try:
        row = db.execute('SELECT metadata, content FROM artifacts WHERE id=? AND owner=? AND kind=?',
                         (artifact_id, client, kind)).fetchone()
    finally:
        db.close()
    if row is None:
        raise HTTPException(404, f'{kind.capitalize()} not found')
    return json.loads(row['metadata']), row['content']


def workbook(content: bytes):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(i.file_size for i in entries) > MAX_EXPANDED_BYTES:
                raise HTTPException(413, 'Workbook expands beyond the supported limit')
            if any('vbaProject' in i.filename for i in entries):
                raise HTTPException(422, 'Macro-enabled workbooks are not supported')
        wb = load_workbook(BytesIO(content), keep_links=False)
        if sum(s.max_row * s.max_column for s in wb) > MAX_CELLS:
            wb.close()
            raise HTTPException(413, 'Workbook exceeds 250,000 cells')
        return wb
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Invalid or unsupported .xlsx workbook')


def worksheet(wb, name: str):
    if name not in wb.sheetnames:
        raise HTTPException(422, f'Unknown worksheet: {name}')
    return wb[name]


def headers(sheet, row: int) -> dict[str, int]:
    if row > sheet.max_row:
        raise HTTPException(422, 'Header row is outside the worksheet')
    result = {}
    for cell in sheet[row]:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        if not name:
            continue
        if name in result:
            raise HTTPException(422, f'Duplicate header: {name}')
        result[name] = cell.column
    if not result:
        raise HTTPException(422, 'Header row is empty')
    return result


class OutputSpec(BaseModel):
    model_config = ConfigDict(extra='forbid')
    template_id: str
    source_sheet: str
    template_sheet: str
    source_header_row: int = Field(default=1, ge=1, le=10000)
    template_header_row: int = Field(default=1, ge=1, le=10000)
    start_row: int = Field(default=2, ge=2, le=10000)
    # Destination header -> source header. Unmapped template cells are preserved.
    mapping: dict[str, str] = Field(min_length=1, max_length=256)

    @model_validator(mode='after')
    def data_after_header(self):
        if self.start_row <= self.template_header_row:
            raise ValueError('start_row must follow template_header_row')
        return self


class JobRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    source_id: str
    outputs: list[OutputSpec] = Field(min_length=1, max_length=10)
    external_reference: str | None = Field(default=None, max_length=200)


router = APIRouter(prefix='/api/v1', tags=['FinAi integration'])


@router.post('/files/{kind}', status_code=201)
def upload(kind: Literal['source', 'template'], file: UploadFile = File(...),
           client: str = Depends(owner)):
    if not (file.filename or '').lower().endswith('.xlsx'):
        raise HTTPException(422, 'Upload an .xlsx file')
    content = file.file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Upload exceeds 20 MiB')
    wb = workbook(content)
    try:
        sheets = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in wb]
    finally:
        wb.close()
    return save(client, kind, {'filename': Path(file.filename).name, 'sheets': sheets}, content)


@router.get('/files/{kind}/{file_id}')
def inspect_file(kind: Literal['source', 'template'], file_id: str,
                 sheet: str, header_row: int = 1, client: str = Depends(owner)):
    if not 1 <= header_row <= 10000:
        raise HTTPException(422, 'header_row must be between 1 and 10000')
    meta, content = fetch(client, file_id, kind)
    wb = workbook(content)
    try:
        return {**meta, 'headers': list(headers(worksheet(wb, sheet), header_row))}
    finally:
        wb.close()


def render(source, template_content: bytes, spec: OutputSpec):
    template = workbook(template_content)
    try:
        src = worksheet(source, spec.source_sheet)
        dest = worksheet(template, spec.template_sheet)
        src_headers = headers(src, spec.source_header_row)
        dst_headers = headers(dest, spec.template_header_row)
        for target, origin in spec.mapping.items():
            if target not in dst_headers or origin not in src_headers:
                raise HTTPException(422, f'Unknown mapping: {target} <- {origin}')
        rows = []
        for row in src.iter_rows(min_row=spec.source_header_row + 1):
            values = [row[src_headers[name] - 1] for name in spec.mapping.values()]
            if all(c.value is None for c in values):
                continue
            if any(c.data_type == 'f' for c in values):
                raise HTTPException(422, 'Mapped source formulas must be converted to values by FinAi')
            rows.append(values)
        if not rows:
            raise HTTPException(422, 'No mapped source data rows')
        if (spec.start_row + len(rows) - 1) * dest.max_column > MAX_CELLS:
            raise HTTPException(413, 'Output exceeds 250,000 cells')
        for offset, values in enumerate(rows):
            for target, source_cell in zip(spec.mapping, values):
                col = dst_headers[target]
                cell = dest.cell(spec.start_row + offset, col)
                if cell.coordinate in dest.merged_cells:
                    raise HTTPException(422, 'Mapped output cells must not be merged')
                if cell.data_type == 'f':
                    raise HTTPException(422, f'Mapping would overwrite template formula {cell.coordinate}')
                style = dest.cell(spec.start_row, col)
                if offset and not cell.has_style:
                    cell._style = copy(style._style)
                value = source_cell.value
                if isinstance(value, str) and ILLEGAL_CHARACTERS_RE.search(value):
                    raise HTTPException(422, 'Source contains unsupported control characters')
                cell.value = value
                # Treat strings such as =SUM(...) as literal untrusted data.
                if isinstance(value, str):
                    cell.data_type = 's'
        buffer = BytesIO()
        template.save(buffer)
        return buffer.getvalue(), len(rows)
    finally:
        template.close()


@router.post('/jobs', status_code=201)
def create_job(request: JobRequest, client: str = Depends(owner)):
    _, source_content = fetch(client, request.source_id, 'source')
    source = workbook(source_content)
    rendered = []
    try:
        for spec in request.outputs:
            _, template_content = fetch(client, spec.template_id, 'template')
            content, count = render(source, template_content, spec)
            rendered.append((spec, content, count))
    finally:
        source.close()
    # Persist the entire completed job atomically; validation failures create nothing.
    job_id = uuid4().hex
    outputs = []
    db = database()
    try:
        with db:
            for index, (spec, content, count) in enumerate(rendered, 1):
                output_id = uuid4().hex
                meta = {'id': output_id, 'filename': f'pexl-{job_id}-{index}.xlsx',
                        'template_id': spec.template_id, 'rows_written': count,
                        'download_path': f'/api/v1/outputs/{output_id}'}
                db.execute('INSERT INTO artifacts VALUES (?, ?, ?, ?, ?)',
                           (output_id, client, 'output', json.dumps(meta), content))
                outputs.append(meta)
            result = {'id': job_id, 'status': 'completed', 'source_id': request.source_id,
                      'external_reference': request.external_reference,
                      'created_at': datetime.now(timezone.utc).isoformat(), 'outputs': outputs}
            db.execute('INSERT INTO artifacts VALUES (?, ?, ?, ?, ?)',
                       (job_id, client, 'job', json.dumps(result), b''))
    finally:
        db.close()
    return result


@router.get('/jobs/{job_id}')
def get_job(job_id: str, client: str = Depends(owner)):
    return fetch(client, job_id, 'job')[0]


@router.get('/outputs/{output_id}')
def download(output_id: str, client: str = Depends(owner)):
    meta, content = fetch(client, output_id, 'output')
    return Response(content, media_type=XLSX_TYPE,
                    headers={'Content-Disposition': f'attachment; filename="{meta["filename"]}"'})


# ---------------------------------------------------------------------------
# Export by document type. The workbook is built by the same TypeScript
# builder as the browser Export button (src/lib/excel-export.ts), bundled by
# `npm run build:export` and run here under Node. Python only authenticates,
# validates the request shape and stores the result.

def export_command(command: str) -> list[str]:
    node = os.environ.get('PEXL_NODE_BIN') or shutil.which('node')
    script = Path(os.environ.get('PEXL_EXPORT_SCRIPT',
                                 Path(__file__).resolve().parent.parent / 'dist-server' / 'export-workbook.mjs'))
    if not node or not script.is_file():
        raise HTTPException(503, 'Workbook export runtime is not available')
    return [node, str(script), command]


def run_export(command: str, payload: dict | None = None) -> dict | list:
    try:
        proc = subprocess.run(export_command(command), capture_output=True, timeout=EXPORT_TIMEOUT_SECONDS,
                              input=json.dumps(payload).encode() if payload is not None else b'')
    except subprocess.TimeoutExpired:
        raise HTTPException(503, 'Workbook export timed out')
    except OSError:
        raise HTTPException(503, 'Workbook export runtime is not available')
    message = proc.stderr.decode(errors='replace').strip()
    if proc.returncode == 2:
        raise HTTPException(422, message or 'Invalid export request')
    if proc.returncode != 0:
        log.error('Export runtime failed (%s): %s', proc.returncode, message)
        raise HTTPException(500, 'Workbook export failed')
    return json.loads(proc.stdout)


@lru_cache(maxsize=4)
def _doc_types(command: tuple[str, ...], mtime: float) -> tuple[dict, ...]:
    return tuple(run_export('doc-types'))


def doc_types() -> list[dict]:
    """Pexl's document types and field keys, read from the TypeScript source of truth."""
    command = export_command('doc-types')
    return list(_doc_types(tuple(command), Path(command[1]).stat().st_mtime))


class ExportRow(BaseModel):
    # FinAi may send amounts as JSON numbers; store them as the string the
    # browser would hold.
    model_config = ConfigDict(extra='forbid', coerce_numbers_to_str=True)
    page: int = Field(ge=1, le=100_000)
    field: str = Field(min_length=1, max_length=200)
    value: str | None = Field(default=None, max_length=10_000)
    confidence: Literal['high', 'medium', 'low'] = 'high'
    wasOcr: bool = False


class ExportFile(BaseModel):
    model_config = ConfigDict(extra='forbid')
    filename: str = Field(min_length=1, max_length=255)
    folder_name: str = Field(default='', max_length=500)
    rows: list[ExportRow]


class ExportOptions(BaseModel):
    model_config = ConfigDict(extra='forbid')
    date_field: Literal['auto', 'billing_date', 'date'] = 'auto'


class ExportRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    doc_type: str
    external_reference: str | None = Field(default=None, max_length=200)
    options: ExportOptions = Field(default_factory=ExportOptions)
    files: list[ExportFile] = Field(min_length=1)


@router.get('/doc-types')
def list_doc_types(client: str = Depends(owner)):
    return doc_types()


@router.post('/exports', status_code=201)
def create_export(request: ExportRequest, client: str = Depends(owner)):
    if len(request.files) > MAX_EXPORT_FILES:
        raise HTTPException(413, f'Export exceeds {MAX_EXPORT_FILES} files')
    if sum(len(f.rows) for f in request.files) > MAX_EXPORT_ROWS:
        raise HTTPException(413, f'Export exceeds {MAX_EXPORT_ROWS:,} rows')
    known = [t['value'] for t in doc_types()]
    if request.doc_type not in known:
        raise HTTPException(422, f'Unknown doc_type: {request.doc_type}. Expected one of: {", ".join(known)}')
    built = run_export('build', request.model_dump())
    content = base64.b64decode(built['xlsx_base64'])
    output_id = uuid4().hex
    meta = {'id': output_id, 'doc_type': request.doc_type, 'filename': built['filename'],
            'download_path': f'/api/v1/outputs/{output_id}',
            'created_at': datetime.now(timezone.utc).isoformat(),
            'external_reference': request.external_reference,
            'unknown_fields': built['unknown_fields']}
    db = database()
    try:
        with db:
            db.execute('INSERT INTO artifacts VALUES (?, ?, ?, ?, ?)',
                       (output_id, client, 'output', json.dumps(meta), content))
    finally:
        db.close()
    return meta


app = FastAPI(title='Pexl Integration API', version='1.0.0')
app.include_router(router)
