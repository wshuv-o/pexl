"""Run: python -m unittest discover -s backend -p test_integration.py"""
from io import BytesIO
import json
import os
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from integration_api import app


def excel(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.active['A2'].font = Font(bold=True)
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {
            'PEXL_API_KEYS': json.dumps({'finai': 'a' * 32, 'other': 'b' * 32}),
            'PEXL_INTEGRATION_DIR': self.temp.name,
        })
        self.env.start()
        self.client = TestClient(app, headers={'Authorization': 'Bearer ' + 'a' * 32})

    def tearDown(self):
        self.client.close()
        self.env.stop()
        self.temp.cleanup()

    def upload(self, kind, rows):
        res = self.client.post('/api/v1/files/' + kind, files={'file': ('test.xlsx', excel(rows))})
        self.assertEqual(res.status_code, 201, res.text)
        return res.json()['id']

    def job(self, source, template):
        return {'source_id': source, 'external_reference': 'finai-123', 'outputs': [{
            'template_id': template, 'source_sheet': 'Sheet', 'template_sheet': 'Sheet',
            'mapping': {'Customer': 'Name', 'Balance': 'Amount'},
        }]}

    def test_multiple_outputs_preserve_styles_types_and_unmapped_formulas(self):
        source = self.upload('source', [['Name', 'Amount'], ['Alice', 12.5], ['Bob', 0]])
        template = self.upload('template', [['Customer', 'Balance', 'Total'], [None, None, '=B2*2']])
        request = self.job(source, template)
        request['outputs'] *= 2
        res = self.client.post('/api/v1/jobs', json=request)
        self.assertEqual(res.status_code, 201, res.text)
        job = res.json()
        self.assertEqual(len(job['outputs']), 2)
        self.assertEqual(self.client.get('/api/v1/jobs/' + job['id']).json(), job)
        data = self.client.get(job['outputs'][0]['download_path'])
        wb = load_workbook(BytesIO(data.content))
        self.assertEqual(wb.active['A2'].value, 'Alice')
        self.assertEqual(wb.active['B2'].value, 12.5)
        self.assertEqual(wb.active['B3'].value, 0)
        self.assertTrue(wb.active['A3'].font.bold)
        self.assertEqual(wb.active['C2'].value, '=B2*2')
        wb.close()
        other = {'Authorization': 'Bearer ' + 'b' * 32}
        self.assertEqual(self.client.get('/api/v1/jobs/' + job['id'], headers=other).status_code, 404)
        self.assertEqual(self.client.get(job['outputs'][0]['download_path'], headers=other).status_code, 404)
        self.assertEqual(self.client.post('/api/v1/jobs', json=request, headers=other).status_code, 404)

    def test_authentication_and_invalid_upload(self):
        self.assertEqual(self.client.get('/api/v1/jobs/x', headers={'Authorization': ''}).status_code, 401)
        with patch.dict(os.environ, {'PEXL_API_KEYS': '{}'}):
            self.assertEqual(self.client.get('/api/v1/jobs/x').status_code, 503)
        res = self.client.post('/api/v1/files/source', files={'file': ('bad.xlsx', b'not excel')})
        self.assertEqual(res.status_code, 422)

    def test_bad_mapping_and_formula_input_rejected(self):
        source = self.upload('source', [['Name', 'Amount'], ['Alice', '=1+2']])
        template = self.upload('template', [['Customer', 'Balance']])
        request = self.job(source, template)
        self.assertEqual(self.client.post('/api/v1/jobs', json=request).status_code, 422)
        request['outputs'][0]['mapping'] = {'Missing': 'Name'}
        self.assertEqual(self.client.post('/api/v1/jobs', json=request).status_code, 422)

    def test_duplicate_headers_and_formula_overwrite_rejected(self):
        source = self.upload('source', [['Name', 'Amount'], ['Alice', 2]])
        template = self.upload('template', [['Customer', 'Balance'], ['=1+2', None]])
        self.assertEqual(self.client.post('/api/v1/jobs', json=self.job(source, template)).status_code, 422)
        duplicate = self.upload('source', [['Name', 'Name'], ['a', 'b']])
        self.assertEqual(self.client.get(f'/api/v1/files/source/{duplicate}',
                                        params={'sheet': 'Sheet'}).status_code, 422)

    def test_literal_equals_string_stays_text(self):
        wb = Workbook()
        wb.active.append(['Name', 'Amount'])
        wb.active.append(['=NOT_A_FORMULA', 1])
        wb.active['A2'].data_type = 's'
        buf = BytesIO()
        wb.save(buf)
        wb.close()
        source = self.client.post('/api/v1/files/source', files={'file': ('x.xlsx', buf.getvalue())}).json()['id']
        template = self.upload('template', [['Customer', 'Balance']])
        job = self.client.post('/api/v1/jobs', json=self.job(source, template)).json()
        output = load_workbook(BytesIO(self.client.get(job['outputs'][0]['download_path']).content))
        self.assertEqual(output.active['A2'].data_type, 's')
        output.close()


if __name__ == '__main__':
    unittest.main()
