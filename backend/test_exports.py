"""Export-by-doc-type API. Needs Node and the export runtime (npm run build:export).

Run: python -m unittest discover -s backend -p test_exports.py
"""
import base64
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import integration_api
from integration_api import app

REPO = Path(__file__).resolve().parent.parent
RUNTIME = REPO / 'dist-server' / 'export-workbook.mjs'
FIXTURES = json.loads((REPO / 'src' / 'test' / 'fixtures' / 'export-requests.json').read_text(encoding='utf-8'))
KEY, OTHER_KEY = 'a' * 32, 'b' * 32


def setUpModule():
    if not shutil.which('node'):
        raise unittest.SkipTest('Node.js is required for the export runtime')
    if not RUNTIME.is_file():
        npm = shutil.which('npm')
        if not npm:
            raise unittest.SkipTest('npm is required to build the export runtime')
        subprocess.run([npm, 'run', 'build:export'], cwd=REPO, check=True, capture_output=True)


def browser_workbook(request):
    """Reference: the browser builder fed rows shaped as the Export button holds them."""
    rows = [{'page': r['page'], 'field': r['field'], 'value': r.get('value'),
             'confidence': 'high', 'wasOcr': False, 'filename': f['filename'],
             **({'folderName': f['folder_name']} if f.get('folder_name') else {}),
             'sessionId': f'session-{i}'}
            for i, f in enumerate(request['files']) for r in f['rows']]
    payload = {'rows': rows, 'provider': 'x', 'options': {
        'forceDocType': request['doc_type'],
        'dateField': request.get('options', {}).get('date_field', 'auto')}}
    proc = subprocess.run(['node', str(RUNTIME), 'build-rows'], input=json.dumps(payload).encode(),
                          capture_output=True, check=True)
    return base64.b64decode(json.loads(proc.stdout)['xlsx_base64'])


def describe(content):
    """Sheet names, values, formulas, merges and styling of every cell."""
    wb = load_workbook(BytesIO(content))
    try:
        return [{
            'name': ws.title,
            'merges': sorted(str(r) for r in ws.merged_cells.ranges),
            'freeze': ws.freeze_panes,
            'cells': {c.coordinate: (c.value, c.number_format, repr(c.font), repr(c.fill),
                                     repr(c.border), repr(c.alignment))
                      for row in ws.iter_rows() for c in row},
        } for ws in wb.worksheets]
    finally:
        wb.close()


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {
            'PEXL_API_KEYS': json.dumps({'finai': KEY, 'other': OTHER_KEY}),
            'PEXL_INTEGRATION_DIR': self.temp.name,
        })
        self.env.start()
        self.client = TestClient(app, headers={'Authorization': 'Bearer ' + KEY})

    def tearDown(self):
        self.client.close()
        self.env.stop()
        self.temp.cleanup()

    def export(self, request):
        res = self.client.post('/api/v1/exports', json=request)
        self.assertEqual(res.status_code, 201, res.text)
        return res.json()

    def test_each_doc_type_matches_browser_export(self):
        for name, request in FIXTURES.items():
            with self.subTest(name):
                created = self.export(request)
                self.assertEqual(created['doc_type'], request['doc_type'])
                self.assertEqual(created['external_reference'], request['external_reference'])
                self.assertEqual(created['download_path'], f"/api/v1/outputs/{created['id']}")
                self.assertTrue(created['filename'].endswith('.xlsx'))
                res = self.client.get(created['download_path'])
                self.assertEqual(res.status_code, 200)
                self.assertIn(created['filename'], res.headers['content-disposition'])
                api, browser = describe(res.content), describe(browser_workbook(request))
                self.assertEqual([s['name'] for s in api], [s['name'] for s in browser])
                self.assertEqual(api, browser)

    def test_formulas_and_merges_survive(self):
        created = self.export(FIXTURES['utility_bill'])
        wb = load_workbook(BytesIO(self.client.get(created['download_path']).content))
        recon = wb['Utility Recon']
        self.assertIn('B5:C5', [str(r) for r in recon.merged_cells.ranges])
        formulas = [c.value for row in recon.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith('=')]
        self.assertTrue(any(f.startswith('=SUM(') for f in formulas))
        self.assertTrue(any('(1.25+2.25)' in f for f in formulas))
        wb.close()

    def test_monthly_bills_bucket_by_date_field(self):
        def request(mode, fields=('billing_date', 'date')):
            rows = []
            for m in range(1, 13):
                rows += [{'page': m, 'field': 'provider_name', 'value': 'ConEd'},
                         {'page': m, 'field': 'total_electricity_bill', 'value': m * 10}]
                rows += [{'page': m, 'field': f, 'value': f'{m:02d}/05/2026'} for f in fields]
            return {'doc_type': 'utility_bill', 'options': {'date_field': mode},
                    'files': [{'filename': 'year.pdf', 'rows': rows}]}

        def months(req):
            wb = load_workbook(BytesIO(self.client.get(self.export(req)['download_path']).content))
            ws = wb['Utility Recon']
            data_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 3).value == 'ConEd')
            out = {ws.cell(6, c).value: ws.cell(data_row, c).value
                   for c in range(5, ws.max_column + 1)
                   if isinstance(ws.cell(6, c).value, str) and ws.cell(6, c).value[:3].isalpha()
                   and ws.cell(6, c).value not in ('Total', 'Comments')}
            wb.close()
            return out

        labels = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
        as_named = {f'{m}-26': (i + 1) * 10 for i, m in enumerate(labels)}
        shifted = {('Dec-25' if i == 0 else f'{labels[i - 1]}-26'): (i + 1) * 10 for i in range(12)}
        self.assertEqual(months(request('billing_date')), shifted)
        self.assertEqual(months(request('date')), as_named)
        self.assertEqual(months(request('auto')), shifted)
        self.assertEqual(months(request('auto', fields=('date',))), as_named)

    def test_doc_types_lists_pexl_types_and_fields(self):
        res = self.client.get('/api/v1/doc-types')
        self.assertEqual(res.status_code, 200)
        types = {t['value']: t for t in res.json()}
        self.assertEqual(set(types), {'utility_bill', 'bank_statement', 'appraisal',
                                      'lease_contract', 'tax', 'others'})
        self.assertEqual(types['utility_bill']['label'], 'Utility Bill')
        fields = {f['value'] for f in types['utility_bill']['fields']}
        self.assertTrue({'total_gas_bill', 'total_electricity_bill', 'billing_date', 'provider_name'} <= fields)

    def test_unknown_fields_are_reported(self):
        self.assertEqual(self.export(FIXTURES['appraisal'])['unknown_fields'], ['total_gas_bill'])

    def test_validation_errors(self):
        request = json.loads(json.dumps(FIXTURES['tax']))
        res = self.client.post('/api/v1/exports', json={**request, 'doc_type': 'invoice'})
        self.assertEqual(res.status_code, 422)
        self.assertIn('Unknown doc_type: invoice', res.json()['detail'])
        bad_rows = [{'page': 0, 'field': 'tax_year', 'value': '2026'}, {'page': 1, 'value': 'x'},
                    {'page': 1, 'field': 'tax_year', 'value': ['list']}, 'not-a-row']
        for row in bad_rows:
            with self.subTest(row=row):
                body = {**request, 'files': [{'filename': 'a.pdf', 'rows': [row]}]}
                self.assertEqual(self.client.post('/api/v1/exports', json=body).status_code, 422)
        self.assertEqual(self.client.post('/api/v1/exports', json={**request, 'files': []}).status_code, 422)
        body = {**request, 'options': {'date_field': 'due_date'}}
        self.assertEqual(self.client.post('/api/v1/exports', json=body).status_code, 422)

    def test_size_limits(self):
        with patch.object(integration_api, 'MAX_EXPORT_ROWS', 3):
            self.assertEqual(self.client.post('/api/v1/exports', json=FIXTURES['tax']).status_code, 413)

    def test_authentication_and_isolation(self):
        for headers in ({'Authorization': ''}, {'Authorization': 'Bearer ' + 'c' * 32}):
            with self.subTest(headers=headers):
                self.assertEqual(self.client.post('/api/v1/exports', json=FIXTURES['tax'],
                                                  headers=headers).status_code, 401)
                self.assertEqual(self.client.get('/api/v1/doc-types', headers=headers).status_code, 401)
        created = self.export(FIXTURES['tax'])
        other = {'Authorization': 'Bearer ' + OTHER_KEY}
        self.assertEqual(self.client.get(created['download_path'], headers=other).status_code, 404)

    def test_missing_runtime_is_503(self):
        with patch.dict(os.environ, {'PEXL_EXPORT_SCRIPT': str(Path(self.temp.name) / 'missing.mjs')}):
            self.assertEqual(self.client.post('/api/v1/exports', json=FIXTURES['tax']).status_code, 503)
            self.assertEqual(self.client.get('/api/v1/doc-types').status_code, 503)
        with patch.dict(os.environ, {'PEXL_API_KEYS': '{}'}):
            self.assertEqual(self.client.post('/api/v1/exports', json=FIXTURES['tax']).status_code, 503)


if __name__ == '__main__':
    unittest.main()
